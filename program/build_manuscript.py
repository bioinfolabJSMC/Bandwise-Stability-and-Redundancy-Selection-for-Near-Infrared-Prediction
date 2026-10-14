import sys,json,re,copy,zipfile,hashlib
from pathlib import Path
import numpy as np,pandas as pd
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent;OUT=ROOT/'results/documents';OUT.mkdir(parents=True,exist_ok=True);A=ROOT/'results/analysis';F=ROOT/'results/figures';TEMPLATE=ROOT/'document_assets/foods_template.docx'
content=json.loads((ROOT/'document_assets/manuscript_content.json').read_text(encoding='utf-8'))
metadata=json.loads((ROOT/'document_assets/article_metadata.json').read_text(encoding='utf-8'))
refs=json.loads((ROOT/'document_assets/references.json').read_text(encoding='utf-8'))
if isinstance(refs,dict):refs=refs.get('references',refs.get('items'))
run_info=json.loads((A/'run_summary.json').read_text(encoding='utf-8'))
if not content.get('analysis_signature') or content['analysis_signature'] != run_info.get('analysis_signature'):
 raise ValueError('The manuscript must be reviewed against the current SPXY analysis signature before generation.')
if content.get('input_sha256') != run_info['input_audit']['sha256']:
 raise ValueError('The manuscript input hash does not match the current 141-variable dataset.')
if re.search(r'\[\[[A-Z_]+\]\]',json.dumps(content,ensure_ascii=False)):
 raise ValueError('An unresolved scientific result placeholder remains in the manuscript.')
refdict={r['key']:r for r in refs};order=list(dict.fromkeys(re.findall(r'@([a-z0-9_]+)',str(content))));missing=set(order)-set(refdict)
if missing:raise ValueError('Missing reference keys '+str(missing))
refnum={k:i+1 for i,k in enumerate(order)}

def format_citation_numbers(numbers):
 numbers=sorted(set(numbers));parts=[];j=0
 while j<len(numbers):
  k=j
  while k+1<len(numbers) and numbers[k+1]==numbers[k]+1:k+=1
  if k-j>=2:parts.append(str(numbers[j])+'–'+str(numbers[k]))
  else:parts.extend(str(v) for v in numbers[j:k+1])
  j=k+1
 return '['+','.join(parts)+']'

def citations(text):
 def repl(m):
  kk=re.findall(r'@([a-z0-9_]+)',m.group());return format_citation_numbers([refnum[k] for k in kk])
 return re.sub(r'\[(?:@[a-z0-9_]+(?:;\s*)?)+\]',repl,text)

ST={'body':'MDPI_3.1_text','plain':'MDPI_3.2_text_no_indent','h1':'MDPI_2.1_heading1','h2':'MDPI_2.2_heading2','caption':'MDPI_5.1_figure_caption','tcaption':'MDPI_4.1_table_caption','figure':'MDPI_5.2_figure','reference':'MDPI_8.1_references'}

def base_document(supp=False):
 d=Document(TEMPLATE);body=d._element.body;sidebar=copy.deepcopy(d.tables[0]._tbl);sec=copy.deepcopy(body.sectPr)
 for child in list(body):body.remove(child)
 body.append(sec)
 if not supp:
  body.insert(0,sidebar)
  # Editable publishing fields only.
  texts=['Academic Editor: [To be assigned]','Received: [To be assigned]','Revised: [To be assigned]','Accepted: [To be assigned]','Published: [To be assigned]','Citation: To be assigned by the publisher.','Copyright: © 2026 by the authors. Submitted for possible open access publication under the terms of the Creative Commons Attribution (CC BY) license.']
  tab=d.tables[0]
  for p,text in zip([p for row in tab.rows for c in row.cells for p in c.paragraphs],texts):
   if p.runs:p.runs[0].text=text
   else:p.add_run(text)
   for run in p.runs[1:]:run.text=''
 else:
  # Supplement starts with the same title/content styles and no editorial sidebar.
  d.sections[0].different_first_page_header_footer=True
 for part in [d.sections[0].header.part,d.sections[0].first_page_header.part,d.sections[0].footer.part,d.sections[0].first_page_footer.part]:
  for text in part.element.iter(qn('w:t')):
   if text.text:
    text.text=text.text.replace('2025','2026').replace('https://doi.org/10.3390/xxxxx','[DOI to be assigned]')
    if text.text.strip()=='14':text.text=text.text.replace('14','x')
 d.core_properties.title=content['title']+(' Supplementary Material' if supp else '')
 d.core_properties.author='; '.join(a['name'] for a in metadata['authors']);d.core_properties.subject='Baishouwu NIR chemometrics';d.core_properties.comments=''
 return d

def append_text(p,text):
 """Use editable Word superscripts/subscripts for inline scientific notation."""
 text=text.replace('mg g−1','mg\u00a0g−1')
 pattern=r'(Cynanchum auriculatum|Cynanchum bungei|10−(?:14|13|12|6)|g−1|\b(?:b_bj|r_j|s_j|q_j|qbar_j|d_j|p_j|y_i|yhat_i|ybar|mean_b|sd_b)\b|F\+)'
 pattern=r'(\[[0-9,–]+\])|'+r'(Xbar_cal,S|ybar_cal|φ_cal|β_j\^\(b\)|ŷ_i\^\(f\)|p_j[+−]|qbar_1|[μσβcz]_j(?:k)?|z_k|q_[123]|[nWPc]_b|V_f|ℒ_K)|'+pattern
 from math_equations import inline_for_token
 position=0
 for match in re.finditer(pattern,text):
  if match.start()>position:p.add_run(text[position:match.start()])
  token=match.group()
  inline=inline_for_token(token) if '_' in token else None
  if token.startswith('['):
   from docx.shared import RGBColor
   run=p.add_run(token);run.font.color.rgb=RGBColor.from_string('075A78')
  elif inline is not None:p._p.append(inline)
  elif token.startswith('Cynanchum'):p.add_run(token).italic=True
  elif token in ['g−1','10−14','10−13','10−12','10−6']:
   base,exp=token.split('−');p.add_run(base);p.add_run('−'+exp).font.superscript=True
  elif token=='F+':p.add_run('F');p.add_run('+').font.superscript=True
  elif token=='ybar':p.add_run('ȳ')
  else:
   base,sub=token.split('_');base={'qbar':'q̄','yhat':'ŷ'}.get(base,base);p.add_run(base);p.add_run(sub).font.subscript=True
  position=match.end()
 if position<len(text):p.add_run(text[position:])
 return p

def para(d,text,role='body'):
 match=re.fullmatch(r'\[EQ:(\d+)\]',text)
 if match:
  from math_equations import add_equation
  return add_equation(d,int(match.group(1)),ST['body'])
 p=d.add_paragraph(style=ST.get(role,role));append_text(p,citations(text));p.paragraph_format.widow_control=True
 if role in ['h1','h2','tcaption']:p.paragraph_format.keep_with_next=True
 if role in ['caption','tcaption']:p.paragraph_format.keep_together=True
 if role=='MDPI_1.2_title':
  nohyphen=OxmlElement('w:suppressAutoHyphens');nohyphen.set(qn('w:val'),'1');p._p.get_or_add_pPr().append(nohyphen)
 return p

def figure(d,stem,label,caption=None):
 entry=next(v for v in json.loads((F/'figure_captions.json').read_text(encoding='utf-8')) if v['stem']==stem)
 p=para(d,'','figure');p.paragraph_format.left_indent=Inches(2608/1440);p.paragraph_format.first_line_indent=Inches(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 width=5.0 if stem=='Figure_1' else 5.45
 shape=p.add_run().add_picture(str(F/(stem+'.png')),width=Inches(width));shape._inline.docPr.set('descr',caption or entry['caption'])
 cap=caption or entry['caption'];cap=cap.replace('block-smoothed stability and redundancy selector','bandwise stability and redundancy selection method')
 p=para(d,f'{label}. {cap}','caption');p.paragraph_format.left_indent=Inches(2608/1440);p.paragraph_format.first_line_indent=Inches(0)
 if p.runs:
  for run in list(p.runs):run._r.getparent().remove(run._r)
  p.add_run(label+'. ').bold=True;append_text(p,cap)
 return p

def table(d,label,caption,headers,rows,widths=None,wide=False,footer=None,row_padding=35):
 p=para(d,f'{label}. {caption}','tcaption');p.paragraph_format.left_indent=Inches(0 if wide else 2608/1440)
 if p.runs:
  for run in list(p.runs):run._r.getparent().remove(run._r)
  p.add_run(label+'. ').bold=True;append_text(p,caption)
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.LEFT
 try:t.style='MDPI_4.1_three_line_table'
 except KeyError:pass
 pr=t._tbl.tblPr
 ind=OxmlElement('w:tblInd');ind.set(qn('w:w'),'0' if wide else '2608');ind.set(qn('w:type'),'dxa');pr.append(ind)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','left','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');borders.append(e)
 old=pr.find(qn('w:tblBorders'))
 if old is not None:pr.remove(old)
 pr.append(borders);total=7.26 if wide else 5.45
 widths=widths or [total/len(headers)]*len(headers);scale=total/sum(widths);widths=[v*scale for v in widths]
 for col,width in zip(t.columns,widths):col.width=Inches(width)
 for j,(cell,text) in enumerate(zip(t.rows[0].cells,headers)):
  cell.width=Inches(widths[j]);cell.text=str(text);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  tcpr=cell._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');bd.append(e);tcpr.append(bd)
 hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
 for row in rows:
  cells=t.add_row().cells
  for j,(cell,text) in enumerate(zip(cells,row)):cell.width=Inches(widths[j]);cell.text=str(text);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  for j,c in enumerate(row.cells):
   tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for edge,value in [('top',row_padding),('bottom',row_padding),('left',35),('right',35)]:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
   for p in c.paragraphs:
    text=p.text
    for run in list(p.runs):run._r.getparent().remove(run._r)
    append_text(p,text)
    p.style='MDPI_4.2_table_body';p.paragraph_format.left_indent=Inches(0);p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==1 and len(headers)>6 else WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
 if footer:para(d,footer,'MDPI_4.3_table_footer')
 return t

def num(v,n=3):return f'{v:.{n}f}'
methods=['Full','UVE','MC-UVE','RT','C-value','BSRS']
names={'Full':'Full-spectrum PLS','UVE':'UVE rank-PLS','MC-UVE':'MC-UVE rank-PLS','RT':'RT-PLS','C-value':'C-value-PLS','BSRS':'BSRS-PLS'}
metrics=pd.read_csv(A/'variable_selection.csv');pre=pd.read_csv(A/'preprocessing.csv');splits=pd.read_csv(A/'split_summary.csv');ab=pd.read_csv(A/'ablation.csv')
chosen=run_info.get('chosen_preprocessing',run_info.get('primary_preprocessing'))
if chosen is None:
 chosen=str(pre.loc[pre.RMSECV.idxmin(),'pretreatment'])
pre_names={'Raw':'raw spectra','MSC':'multiplicative scatter correction','SNV':'standard normal variate','CWT':'continuous wavelet transform','D1':'first derivative','D2':'second derivative'}
chosen_name=pre_names.get(chosen,str(chosen))

def main_component(d,kind):
 if kind=='data':
  v=json.loads((A/'input_data.json').read_text(encoding='utf-8'))
  table(d,'Table 1','Structure of the near-infrared spectral dataset.',
   ['Samples','Spectral variables','Wavelength range (nm)'],
   [[v.get('n',v.get('samples')),v['variables'],f"{v['range_nm'][0]:.2f}–{v['range_nm'][1]:.2f}"]],widths=[.9,1.5,3.05])
  figure(d,'Figure_2','Figure 2')
 if kind=='split':
  table(d,'Table 2','Reference-content distributions in the SPXY calibration and prediction sets. All values except n are in mg g−1.',
   ['Set','n','Minimum','Maximum','Mean','SD'],
   [[r['set'],int(r.n),num(r['min']),num(r['max']),num(r['mean']),num(r.sd)] for _,r in splits.iterrows()],
   widths=[1.45,.4,.9,.9,.9,.9],wide=False)
  figure(d,'Figure_3','Figure 3')
 if kind=='pre':
  table(d,'Table 3','Full-spectrum PLS under each preprocessing condition. Model selection used calibration tuning RMSECV.',
   ['Preprocessing','LV','RMSECV','RMSEP','R²p','RPD'],
   [[r.pretreatment,int(r.LV),num(r.RMSECV),num(r.RMSEP),num(r.R2p),num(r.RPD)] for _,r in pre.iterrows()],
   widths=[1.35,.4,1.0,1.0,.8,.8],footer='Errors are in mg g−1. D1 and D2 denote first and second derivatives. Prediction errors did not determine preprocessing selection.')
  figure(d,'Figure_4','Figure 4')
 if kind=='selection':
  table(d,'Table 4',f'Variable-selection comparisons using {chosen_name}, selected from calibration data.',
   ['Method','Variables','LV','RMSECV','RMSEP','R²p'],
   [[names[r.method],int(r.variables),int(r.LV),num(r.RMSECV),num(r.RMSEP),num(r.R2p)] for _,r in metrics.iterrows()],
   widths=[1.65,.75,.35,.85,.85,.7],footer='Errors are in mg g−1. LV is the effective final component count. UVE rank and MC-UVE rank use CV-tuned reliability prefixes; C-value-PLS uses the adapted random inclusion design.')
  figure(d,'Figure_5','Figure 5');figure(d,'Figure_6','Figure 6')
 if kind=='bsrs':
  figure(d,'Figure_7','Figure 7')
  r=metrics[metrics.method=='BSRS'].iloc[0]
  rows=[['Complete BSRS',int(r.variables),int(r.LV),num(r.RMSECV),num(r.RMSEP)]]
  labels={'BSRS-nosign':'Without sign weighting','BSRS-nosmooth':'Without smoothing','BSRS-noredund':'Without redundancy penalty'}
  rows.extend([[labels[r.variant],int(r.variables),int(r.LV),num(r.RMSECV),num(r.RMSEP)] for _,r in ab.iterrows()])
  table(d,'Table 5','Component ablation of BSRS, with calibration tuning repeated for each variant.',
   ['Variant','Variables','LV','RMSECV','RMSEP'],rows,widths=[2.2,.75,.35,.7,.8],wide=False,
   footer='Errors are in mg g−1. Ablations assess the roles of the three algorithm components; the complete method was specified before prediction-set evaluation.')

def reference(d,ref,n):
 p=para(d,'','reference');p.paragraph_format.left_indent=Inches(425/1440);p.paragraph_format.first_line_indent=Inches(-425/1440)
 # Explicit numbers suppress inherited list numbering and avoid double labels.
 pr=p._p.get_or_add_pPr();nppr=OxmlElement('w:numPr');nid=OxmlElement('w:numId');nid.set(qn('w:val'),'0');nppr.append(nid);pr.append(nppr)
 authors=ref.get('authors',[])
 if authors and isinstance(authors[0],dict):
  def format_author(a):
   given=a.get('given','');initials=''.join('-'.join((part[0]+'.') if part else '' for part in token.split('-')) for token in given.split())
   return a.get('family','')+', '+initials
  author='; '.join(a.get('formatted',format_author(a)) for a in authors)
 else:author='; '.join(authors)
 p.add_run(f'{n}. {author}'+(' ' if author.endswith('.') else '. ')+f'{ref["title"].rstrip(".")}. ')
 p.add_run(ref.get('journal_abbrev',ref.get('journal_abbreviation',ref.get('journal','')))).italic=True;p.add_run(' ');p.add_run(str(ref['year'])).bold=True;p.add_run(', ')
 if ref.get('volume'):p.add_run(str(ref['volume'])).italic=True;p.add_run(', ')
 p.add_run(str(ref.get('pages',ref.get('article_number',''))).replace('--','–')+'. ')
 if ref.get('doi'):
  if ref['key'] in ['cwt','spectral_physics']:p.add_run().add_break()
  p.add_run('https://doi.org/'+ref['doi'])
 elif ref.get('url'):p.add_run(ref['url'])
 return p

def author_block(d):
 p=para(d,'','MDPI_1.3_authornames')
 for j,author in enumerate(metadata['authors']):
  if j:p.add_run(' and ' if j==len(metadata['authors'])-1 else ', ')
  p.add_run(author['name'])
  marker=','.join(str(v) for v in author['affiliations'])+(',*' if author['corresponding'] else '')
  p.add_run(' '+marker).font.superscript=True
 for aff in metadata['affiliations']:
  p=para(d,'','MDPI_1.6_affiliation')
  p.paragraph_format.left_indent=Inches(2608/1440+.18)
  p.paragraph_format.first_line_indent=Inches(-.18)
  p.paragraph_format.tab_stops.add_tab_stop(Inches(2608/1440+.18))
  p.paragraph_format.space_after=Pt(2)
  p.add_run(str(aff['number'])).font.superscript=True
  p.add_run('\t'+aff['text'].replace('300071, China','300071,\u00a0China'))
  for run in p.runs:run.font.size=Pt(8)
 correspondence=metadata['correspondence_line']
 for person in metadata['correspondence']:correspondence=correspondence.replace(' ('+person['initials']+')','\u00a0('+person['initials']+')')
 correspondence=correspondence.replace('; '+metadata['correspondence'][-1]['email'],';\n'+metadata['correspondence'][-1]['email'])
 p=para(d,correspondence,'MDPI_1.6_affiliation')
 nohyphen=OxmlElement('w:suppressAutoHyphens');nohyphen.set(qn('w:val'),'1');p._p.get_or_add_pPr().append(nohyphen)
 p.paragraph_format.space_after=Pt(6)
 for run in p.runs:run.font.size=Pt(8)

def abbreviations(d):
 para(d,'Abbreviations','h1')
 values=[('NIR','Near-infrared'),('PLS','Partial least squares'),('BSRS','Bandwise stability and redundancy selection'),
  ('UVE','Uninformative variable elimination'),('MC-UVE','Monte Carlo uninformative variable elimination'),
  ('RT','Randomization test'),('MSC','Multiplicative scatter correction'),('SNV','Standard normal variate'),
  ('CWT','Continuous wavelet transform'),('LV','Latent variable'),
  ('SPXY','Sample set partitioning based on joint x–y distances'),('RMSECV','Root mean square error of cross validation'),
  ('RMSEP','Root mean square error of prediction'),('RPD','Ratio of performance to deviation')]
 for name,value in values:
  p=para(d,'','plain');p.paragraph_format.tab_stops.add_tab_stop(Inches(2608/1440+1.0))
  p.add_run(name+'\t'+value)
  p.paragraph_format.space_after=Pt(0)
  for run in p.runs:run.font.size=Pt(9)

def main_document():
 d=base_document();para(d,'Article','MDPI_1.1_article_type')
 para(d,content['title'],'MDPI_1.2_title');author_block(d)
 p=para(d,'Abstract','h1');p.paragraph_format.space_before=Pt(6)
 p=para(d,content['abstract'],'MDPI_1.7_abstract');p.paragraph_format.space_before=Pt(0)
 p=para(d,'','MDPI_1.8_keywords');p.add_run('Keywords: ').bold=True;append_text(p,content['keywords'])
 para(d,'','MDPI_1.9_line')
 for section in content['sections']:
  para(d,section['heading'],'h1')
  for text in section.get('paragraphs',[]):para(d,text)
  for sub in section.get('subsections',[]):
   para(d,sub['heading'],'h2')
   for pi,text in enumerate(sub['paragraphs']):
    p=para(d,text)
    if pi+1<len(sub['paragraphs']) and sub['paragraphs'][pi+1].startswith('[EQ:'):
     p.paragraph_format.keep_with_next=True
    if sub['heading'].startswith('2.3.') and pi==0:figure(d,'Figure_1','Figure 1')
    if sub['heading'].startswith('4.1.') and pi==0:main_component(d,'data')
    if sub['heading'].startswith('4.1.') and pi==1:main_component(d,'split')
   heading=sub['heading']
   if heading.startswith('4.2.'):main_component(d,'pre')
   if heading.startswith('4.3.'):main_component(d,'selection')
   if heading.startswith('4.4.'):main_component(d,'bsrs')
 back=[('Supplementary Materials','Figures S1–S4 and Tables S1–S5 provide transformed spectra, prediction residuals, BSRS diagnostics, component ablations, SPXY sample assignments, additional performance metrics, selected wavelengths and computational settings. The accompanying program reproduces the analysis, figures and Word files.'),
  ('Author Contributions','[The authors should complete and approve their individual CRediT roles.]'),
  ('Funding',metadata['funding_statement']),
  ('Institutional Review Board Statement','Not applicable.'),('Informed Consent Statement','Not applicable.'),
  ('Data Availability Statement','The data, numerical results and code used in this study accompany the manuscript. [Add the public repository DOI before submission.]'),
  ('Acknowledgments','AI assistance was used for code development, graphics and manuscript drafting. [The authors should verify the analysis and complete the tool/version disclosure.]'),
  ('Conflicts of Interest','[The authors should complete their conflict-of-interest declaration.]')]
 for label,text in back:
  p=para(d,'','MDPI_6.2_back_matter');p.add_run(label+': ').bold=True;append_text(p,text)
 abbreviations(d)
 p=para(d,'References','h1');p.paragraph_format.left_indent=Inches(0)
 for k in order:reference(d,refdict[k],refnum[k])
 save_preserving(d,OUT/'Foods_manuscript.docx')

def supplement():
 d=base_document(True);para(d,'Supplementary Material','MDPI_1.1_article_type');para(d,content['title'],'MDPI_1.2_title');author_block(d);para(d,'Supplementary figures and tables','h1')
 para(d,'The analysis uses 84 samples, each measured at 141 near-infrared spectral variables. SPXY partitions the samples into 63 calibration and 21 prediction samples using normalized Euclidean distances in the full raw spectral matrix and the reference response. The main manuscript defines the preprocessing, variable-selection methods and calibration procedure.')
 for j in range(1,5):figure(d,f'Figure_S{j}',f'Figure S{j}')
 assignment=pd.read_csv(A/'sample_assignment.csv')
 row_column='sample_row' if 'sample_row' in assignment.columns else 'MAT_row'
 table(d,'Table S1','Complete sample assignment for the SPXY 75:25 partition.',
  ['Sample row','Source ID','Reference\n(mg g−1)','Set'],
  [[int(r[row_column]),r.source_sample_id,num(r.reference_mg_g,2),r['set']] for _,r in assignment.iterrows()],
  widths=[.8,1.1,1.7,1.85],footer='Source identifiers link the numerical data to the companion reference-content records. No sample exclusion was applied.',row_padding=25)
 table(d,'Table S2','Additional prediction performance for the SPXY prediction set.',
  ['Method','MAE','Bias','RPD','RPIQ'],
  [[names[r.method],num(r.MAE),num(r.bias),num(r.RPD),num(r.RPIQ)] for _,r in metrics.iterrows()],
  widths=[2.05,.85,.85,.85,.85],footer='MAE and bias are in mg g−1. RPD and RPIQ are dimensionless.')
 table(d,'Table S3','Descriptive selection consistency across calibration CV folds.',
  ['Method','Final variables','Mean fold Jaccard'],
  [[names[r.method],int(r.variables),num(r.mean_fold_Jaccard)] for _,r in metrics.iterrows()],
  widths=[2.0,1.0,2.45],footer='Pairwise Jaccard similarity is averaged across ten pairs of five calibration CV training folds. The subsets may have different sizes.')
 v=json.loads((A/'selection_details.json').read_text(encoding='utf-8'))['BSRS']
 grid={'uniform interpolated':'Uniform grid','original':'Measured grid'}.get(v.get('working_grid'),v.get('working_grid','Measured grid'))
 table(d,'Table S4','Final BSRS wavelength positions.',
  ['Grid','Count','Selected wavelength positions (nm)'],
  [[grid,len(v['selected_wavelengths_nm']),', '.join(f'{q:.1f}' for q in sorted(v['selected_wavelengths_nm']))]],
  widths=[1.0,.5,5.2],wide=True,footer='These positions are predictive features rather than molecular assignments. selected_wavelengths.csv records the selections of every comparator at full precision.')
 settings=[['Input','84 samples; 141 near-infrared spectral variables'],['Partition','SPXY; 63 calibration and 21 prediction samples'],['SPXY distance','Full raw spectral Euclidean distance divided by its maximum plus response Euclidean distance divided by its maximum'],['Random seed','20261006; deterministic offsets for selector and fold draws'],['Preprocessing selection','Minimum five-fold full-spectrum calibration tuning RMSECV'],['PLS','Mean centering; no autoscaling; 1–10 requested LVs'],['Derivative','Uniform grid; cubic Savitzky–Golay; 15 points; mirror edges'],['CWT','Uniform grid; Mexican hat; scale 10; integration precision 10; 40-point reflected boundary extension'],['Stability sampling','100 subsamples; 80% without replacement'],['UVE','Leave-one-out reliability; appended Gaussian noise; CV-tuned rank prefix'],['MC-UVE','100-subsample reliability; appended Gaussian noise; CV-tuned rank prefix'],['RT','199 response permutations; finite-sample correction'],['C-value adapted','2p binary subset rows; 70% inclusion; three-fold internal CV'],['BSRS','Sign-weighted log reliability; five-point smoothing; redundancy weight 2'],['Subset budgets','4, 8, 16, 32, 64, 71 and 128']]
 table(d,'Table S5','Settings required to reproduce the comparisons.',['Setting','Value'],settings,widths=[1.6,3.85])
 para(d,'Computational reproducibility','h1')
 para(d,'The complete program includes data_io.py, analysis.py, plot_figures.py, verify_results.py, predict_new.py and build_manuscript.py. run_all.py executes the pipeline. The input MAT file contains the 141-variable spectral matrix, its wavelength vector and the reference response; its SHA-256 hash, SPXY assignments, model coefficients and numerical outputs are included. Requirements and tested package versions are documented. Figures are generated directly from result tables.')
 save_preserving(d,OUT/'Foods_supplementary_material.docx')

def save_preserving(d,path):
 temporary=OUT/('temporary_'+path.name);d.save(temporary)
 # Restore every opaque preserve-only template part byte for byte.
 preserve=['word/styles.xml','word/numbering.xml','word/theme/theme1.xml','word/fontTable.xml','word/webSettings.xml','word/customizations.xml','word/footnotes.xml','word/endnotes.xml']
 with zipfile.ZipFile(TEMPLATE) as src,zipfile.ZipFile(temporary) as built,zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as dst:
  for item in built.infolist():
   data=src.read(item.filename) if item.filename in preserve and item.filename in src.namelist() else built.read(item.filename)
   if item.filename=='word/settings.xml':
    from lxml import etree
    root=etree.fromstring(data);node=root.find(qn('w:updateFields'))
    if node is None:node=OxmlElement('w:updateFields');root.append(node)
    node.set(qn('w:val'),'true');data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
   dst.writestr(item,data)
 temporary.unlink()
 print('Saved',path,flush=True)

main_document();supplement()
(OUT/'citation_order.json').write_text(json.dumps([{'number':refnum[k],'key':k,'doi':refdict[k].get('doi'),'url':refdict[k].get('url')} for k in order],indent=2),encoding='utf-8')
