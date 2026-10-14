"""Editable native Word equations for the BSRS derivation.

The OMML below is generated directly from our own equation definitions.
No Office stylesheet, font, or other vendor component is redistributed.
"""
from copy import deepcopy
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

def element(tag,*children,**attributes):
    node=OxmlElement('m:'+tag)
    for k,v in attributes.items():node.set(qn('m:'+k),str(v))
    for child in children:
        if isinstance(child,(tuple,list)):
            for c in child:node.append(deepcopy(c))
        elif child is not None:node.append(deepcopy(child))
    return node

def r(text,plain=False):
    node=element('r')
    if plain:node.append(element('rPr',element('sty',val='p')))
    t=element('t');t.text=text
    t.set(qn('xml:space'),'preserve');node.append(t)
    return node

def seq(*nodes):return list(nodes)
def sub(base,index):return element('sSub',element('e',base),element('sub',index))
def sup(base,power):return element('sSup',element('e',base),element('sup',power))
def ss(base,index,power):return element('sSubSup',element('e',base),element('sub',index),element('sup',power))
def frac(num,den):return element('f',element('num',num),element('den',den))
def bar(base):return element('acc',element('accPr',element('chr',val='̅')),element('e',base))
def hat(base):return element('acc',element('accPr',element('chr',val='̂')),element('e',base))
def bracket(base,left='(',right=')'):
    return element('d',element('dPr',element('begChr',val=left),element('endChr',val=right)),element('e',base))
def absolute(base):return bracket(base,'|','|')
def norm(base):return bracket(base,'‖','‖')
def summation(base,lower,upper=None):
    props=[element('chr',val='∑'),element('limLoc',val='undOvr')]
    if upper is None:props.append(element('supHide',val='1'))
    return element('nary',element('naryPr',props),element('sub',lower),element('sup',upper if upper is not None else []),element('e',base))
def root(base):return element('rad',element('radPr',element('degHide',val='1')),element('deg'),element('e',base))
def operator(text):return r(text,True)
def jvar(name):return sub(r(name),r('j'))
def beta():return ss(r('β'),r('j'),bracket(r('b')))
def bvector():return sup(r('β'),bracket(r('b')))
def mu():return jvar('μ')
def sigma():return jvar('σ')
def reliability():return jvar('r')
def sign():return jvar('s')
def importance():return jvar('q')
def smooth():return bar(jvar('q'))

def definition(number):
    eq=operator(' = ');plus=operator(' + ');minus=operator(' − ')
    if number==1:
        w=sub(r('W'),r('b'));p=sub(r('P'),r('b'));c=sub(r('c'),r('b'))
        return seq(bvector(),eq,w,sup(bracket(seq(sup(p,r('T')),w)),operator('+')),c)
    if number==2:
        mean=frac(summation(beta(),seq(r('b'),eq,r('1')),r('B')),r('B'))
        variance=frac(summation(sup(bracket(seq(beta(),minus,mu())),r('2')),seq(r('b'),eq,r('1')),r('B')),seq(r('B'),minus,r('1')))
        return seq(mu(),eq,mean,operator(',     '),sup(sigma(),r('2')),eq,variance)
    if number==3:
        return seq(reliability(),eq,frac(absolute(mu()),seq(operator('max'),bracket(seq(sigma(),operator(', '),r('ε'))))),operator(',   '),r('ε'),eq,sup(r('10'),r('−14')))
    if number==4:
        value=frac(summation(seq(operator('sign'),bracket(beta())),seq(r('b'),eq,r('1')),r('B')),r('B'))
        return seq(sign(),eq,absolute(value),eq,absolute(seq(ss(r('p'),r('j'),operator('+')),minus,ss(r('p'),r('j'),operator('−')))))
    if number==5:
        return seq(importance(),eq,operator('log'),bracket(seq(r('1'),plus,reliability())),sign(),operator(',   0 ≤ '),importance(),operator(' ≤ log'),bracket(seq(r('1'),plus,reliability())))
    if number==6:
        qref=sub(r('q'),seq(r('ρ'),bracket(seq(r('j'),plus,r('ℓ')))))
        loss=summation(sup(bracket(seq(qref,minus,r('a'))),r('2')),seq(r('ℓ'),eq,r('−2')),r('2'))
        argmin=element('limLow',element('e',operator('arg min')),element('lim',r('a')))
        average=frac(summation(qref,seq(r('ℓ'),eq,r('−2')),r('2')),r('5'))
        return seq(smooth(),eq,argmin,loss,eq,average)
    if number==7:
        centered=seq(jvar('x'),minus,bar(jvar('x')),r('1'))
        normalized=frac(centered,seq(operator('max'),bracket(seq(norm(centered),operator(', '),sup(r('10'),r('−12'))))))
        return seq(jvar('z'),eq,normalized,operator(',   '),sub(r('c'),r('jk')),eq,sup(jvar('z'),r('T')),sub(r('z'),r('k')))
    if number==8:
        zk=sub(r('z'),r('k'))
        residual=seq(bracket(seq(r('I'),minus,zk,sup(zk,r('T')))),jvar('z'))
        return seq(sup(norm(residual),r('2')),eq,r('1'),minus,sup(sub(r('c'),r('jk')),r('2')))
    if number==9:
        maximum=element('limLow',element('e',operator('max')),element('lim',seq(r('k'),operator(' ∈ '),r('S'))))
        return seq(jvar('d'),bracket(r('S')),eq,maximum,absolute(sub(r('c'),r('jk'))),operator(',   '),jvar('d'),bracket(operator('∅')),eq,r('0'))
    if number==10:
        d=seq(jvar('d'),bracket(r('S')))
        denominator=seq(r('1'),plus,r('λ'),sup(bracket(d),r('2')))
        return seq(jvar('m'),bracket(r('S')),eq,frac(smooth(),denominator),operator(',   '),r('λ'),eq,r('2'))
    if number==11:
        lower=seq(r('j'),operator(' ∉ '),sub(r('S'),seq(r('t'),minus,r('1'))))
        maximum=element('limLow',element('e',operator('arg max')),element('lim',lower))
        previous=sub(r('S'),seq(r('t'),minus,r('1')))
        return seq(sub(r('j'),r('t')),eq,maximum,operator(' '),jvar('m'),bracket(previous),operator(',   '),sub(r('S'),r('t')),eq,previous,operator(' ∪ '),bracket(sub(r('j'),r('t')),'{','}'))
    if number==12:
        argmin=element('limLow',element('e',operator('arg min')),element('lim',seq(r('K'),operator(' ∈ '),r('𝒦'),operator(', '),r('L'),operator(' ∈ '),sub(r('ℒ'),r('K')))))
        yh=ss(hat(r('y')),r('i'),bracket(r('f')))
        residual=sup(bracket(seq(sub(r('y'),r('i')),minus,yh,bracket(seq(r('K'),operator(', '),r('L'))))),r('2'))
        total=summation(summation(residual,seq(r('i'),operator(' ∈ '),sub(r('V'),r('f')))),seq(r('f'),eq,r('1')),r('5'))
        return seq(bracket(seq(hat(r('K')),operator(', '),hat(r('L')))),eq,argmin,root(frac(total,sub(r('n'),operator('cal')))))
    if number==13:
        transformed=sub(r('φ'),operator('cal'))
        centered=seq(sub(bracket(seq(transformed,bracket(r('x')))),r('S')),minus,bar(sub(r('X'),seq(operator('cal'),operator(', '),r('S')))))
        return seq(hat(r('y')),bracket(r('x')),eq,bar(sub(r('y'),operator('cal'))),plus,sup(bracket(centered),r('T')),sub(hat(r('β')),r('S')))
    raise KeyError(number)

def add_equation(document,number,style='MDPI_3.1_text'):
    p=document.add_paragraph(style=style)
    fmt=p.paragraph_format
    fmt.first_line_indent=Inches(0);fmt.left_indent=Inches(2608/1440)
    fmt.space_before=Pt(5);fmt.space_after=Pt(7)
    fmt.keep_together=True;fmt.widow_control=True
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    fmt.tab_stops.add_tab_stop(Inches(4.53),WD_TAB_ALIGNMENT.CENTER)
    fmt.tab_stops.add_tab_stop(Inches(7.26),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t')
    p._p.append(element('oMath',definition(number)))
    p.add_run(f'\t({number})').font.size=Pt(10)
    return p

def inline_for_token(token):
    """Typeset named scientific symbols that occur inside explanatory prose."""
    if token=='β_j^(b)':nodes=beta()
    elif token=='yhat_i':nodes=sub(hat(r('y')),r('i'))
    elif token=='ŷ_i^(f)':nodes=ss(hat(r('y')),r('i'),bracket(r('f')))
    elif token=='Xbar_cal,S':nodes=bar(sub(r('X'),seq(operator('cal'),operator(', '),r('S'))))
    elif token=='ybar_cal':nodes=bar(sub(r('y'),operator('cal')))
    elif token in ['p_j+','p_j−']:nodes=ss(r('p'),r('j'),operator(token[-1]))
    elif token in ['qbar_j','qbar_1']:nodes=bar(sub(r('q'),r(token.split('_')[1])))
    elif '_' in token:
        base,index=token.split('_',1)
        nodes=sub(r(base),operator(index) if index=='cal' else r(index))
    else:return None
    return element('oMath',nodes)
